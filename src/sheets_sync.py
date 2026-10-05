"""
Google Sheets & Dual-Sheet Exporter Module
Syncs data to Google Sheets via gspread or creates an Excel workbook containing:
1. 'Working_Paper_Result'
2. 'Dashboard'
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Dict, Any, List


class DualSheetExporter:
    def __init__(self, excel_path: str):
        self.excel_path = excel_path

    def add_dashboard_sheet(self, metrics: Dict[str, Any], executive_summary_text: str):
        """
        Creates and styles the 'Dashboard' sheet in the Excel workbook.
        """
        wb = openpyxl.load_workbook(self.excel_path)
        
        # If Dashboard already exists, remove it
        if "Dashboard" in wb.sheetnames:
            del wb["Dashboard"]

        # Create Dashboard as first sheet
        ws = wb.create_sheet(title="Dashboard", index=0)
        ws.views.sheetView[0].showGridLines = True

        # Palettes
        navy_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
        card_fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
        accent_blue = PatternFill(start_color="2B6CB0", end_color="2B6CB0", fill_type="solid")
        green_fill = PatternFill(start_color="2F855A", end_color="2F855A", fill_type="solid")
        amber_fill = PatternFill(start_color="DD6B20", end_color="DD6B20", fill_type="solid")

        white_title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        white_subtitle_font = Font(name="Calibri", size=11, color="E2E8F0")
        card_header_font = Font(name="Calibri", size=10, bold=True, color="4A5568")
        card_val_font = Font(name="Calibri", size=16, bold=True, color="1A202C")
        section_font = Font(name="Calibri", size=12, bold=True, color="1A365D")

        thin_border = Border(
            left=Side(style='thin', color='CBD5E0'),
            right=Side(style='thin', color='CBD5E0'),
            top=Side(style='thin', color='CBD5E0'),
            bottom=Side(style='thin', color='CBD5E0')
        )

        # Title Banner (Rows 1-2, Cols A-L)
        ws.merge_cells("A1:L1")
        ws["A1"] = "SOUTHCITY FINANCE & IT AUTOMATION — ADVANCE SETTLEMENT DASHBOARD"
        ws["A1"].font = white_title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws["A1"].fill = navy_fill

        ws.merge_cells("A2:L2")
        ws["A2"] = f"Executive Reconciliation Report | Periode: April 2026 | Total Settlement Rate: {metrics['settlement_rate_pct']}%"
        ws["A2"].font = white_subtitle_font
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws["A2"].fill = navy_fill

        # Fill navy for all banner cells
        for col_i in range(1, 13):
            ws.cell(row=1, column=col_i).fill = navy_fill
            ws.cell(row=2, column=col_i).fill = navy_fill

        ws.row_dimensions[1].height = 36
        ws.row_dimensions[2].height = 22
        ws.row_dimensions[3].height = 10

        # KPI CARDS (Row 4 to 5) — wider merge ranges
        kpis = [
            ("TOTAL ADVANCE", f"Rp {metrics['total_advance']:,.0f}", "B", "D"),
            ("TOTAL REALISASI", f"Rp {metrics['total_realization']:,.0f}", "E", "G"),
            ("SISA SALDO (OUTSTANDING)", f"Rp {metrics['total_saldo']:,.0f}", "H", "J"),
        ]

        for title, val, col_start, col_end in kpis:
            ws.merge_cells(f"{col_start}4:{col_end}4")
            ws[f"{col_start}4"] = title
            ws[f"{col_start}4"].font = card_header_font
            ws[f"{col_start}4"].alignment = Alignment(horizontal="center", vertical="center")
            ws[f"{col_start}4"].fill = card_fill

            ws.merge_cells(f"{col_start}5:{col_end}5")
            ws[f"{col_start}5"] = val
            ws[f"{col_start}5"].font = card_val_font
            ws[f"{col_start}5"].alignment = Alignment(horizontal="center", vertical="center")
            ws[f"{col_start}5"].fill = card_fill

            for col in range(ord(col_start) - 64, ord(col_end) - 64 + 1):
                for r in [4, 5]:
                    ws.cell(r, col).border = thin_border
                    ws.cell(r, col).fill = card_fill

        ws.row_dimensions[4].height = 22
        ws.row_dimensions[5].height = 30

        # STATUS SUMMARY TABLE (Rows 7-11)
        ws.cell(row=7, column=2, value="RINGKASAN STATUS TRANSAKSI").font = section_font
        ws.merge_cells("B7:F7")

        headers = ["Status", "Jumlah Item", "Keterangan"]
        header_cols = [2, 4, 5]
        header_merges = [("B8", "C8"), None, ("E8", "G8")]
        for i, (col_idx, h) in enumerate(zip(header_cols, headers)):
            cell = ws.cell(row=8, column=col_idx, value=h)
            cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            cell.fill = accent_blue
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        # Merge header cells
        ws.merge_cells("B8:C8")
        ws.merge_cells("E8:G8")
        ws.cell(row=8, column=3).fill = accent_blue
        ws.cell(row=8, column=3).border = thin_border
        ws.cell(row=8, column=6).fill = accent_blue
        ws.cell(row=8, column=6).border = thin_border
        ws.cell(row=8, column=7).fill = accent_blue
        ws.cell(row=8, column=7).border = thin_border

        status_rows = [
            ("Fully Settled (Lunas)", metrics['count_settled'], "100% Selesai & Cocok dengan GL"),
            ("Partially Settled (Sebagian)", metrics['count_partial'], "PBB JV 2 Summarecon (Sisa Rp 107.1M)"),
            ("Unsettled (Belum Selesai)", metrics['count_unsettled'], "Tidak ada transaksi tanpa matching"),
        ]

        for idx, (stat, cnt, desc) in enumerate(status_rows, start=9):
            ws.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=3)
            ws.cell(row=idx, column=2, value=stat).border = thin_border
            ws.cell(row=idx, column=3).border = thin_border
            ws.cell(row=idx, column=4, value=cnt).alignment = Alignment(horizontal="center")
            ws.cell(row=idx, column=4).border = thin_border
            ws.merge_cells(start_row=idx, start_column=5, end_row=idx, end_column=7)
            ws.cell(row=idx, column=5, value=desc).border = thin_border
            ws.cell(row=idx, column=6).border = thin_border
            ws.cell(row=idx, column=7).border = thin_border

        # EXECUTIVE SUMMARY TEXT SECTION (Row 13 onwards)
        ws.cell(row=13, column=2, value="EXECUTIVE SUMMARY (LAPORAN REKONSILIASI)").font = section_font
        ws.merge_cells("B13:L13")

        from src.cleaner import bersihkan_markdown_dan_asterik

        # Bersihkan dari semua asterik dan markdown
        clean_text = bersihkan_markdown_dan_asterik(executive_summary_text)

        header_fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
        section_title_font = Font(name="Calibri", size=11, bold=True, color="1A365D")
        body_font = Font(name="Calibri", size=10, color="2D3748")
        bullet_font = Font(name="Calibri", size=10, color="2D3748")

        # Merge Executive Summary text across B-L (much wider)
        current_row = 15
        for line in clean_text.split("\n"):
            line_clean = line.strip()

            # Lewati baris kosong atau garis pembatas teks biasa
            if not line_clean or line_clean.startswith("---") or line_clean.startswith("===") or line_clean.startswith("___") or line_clean.startswith("----"):
                continue

            # Lewati basa-basi chatbot AI
            lower_line = line_clean.lower()
            if any(basa in lower_line for basa in [
                "tentu, berikut", "berikut adalah executive summary", "laporan ini kami sampaikan untuk menjadi",
                "demikian executive summary ini"
            ]):
                continue

            # Cek tipe baris
            is_section_header = (
                (len(line_clean) > 2 and line_clean[0].isdigit() and line_clean[1] in [".", ")"]) or
                line_clean.startswith("EXECUTIVE SUMMARY") or
                line_clean.startswith("Rekonsiliasi Advance Settlement")
            )
            is_bullet = line_clean.startswith("- ") or line_clean.startswith("• ")

            ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=12)
            cell = ws.cell(row=current_row, column=2, value=line_clean)

            if is_section_header:
                cell.font = section_title_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                ws.row_dimensions[current_row].height = 24
                for col_i in range(2, 13):
                    ws.cell(row=current_row, column=col_i).fill = header_fill
                    ws.cell(row=current_row, column=col_i).border = thin_border
            elif is_bullet:
                cell.font = bullet_font
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
                approx_lines = max(1, len(line_clean) // 120 + 1)
                ws.row_dimensions[current_row].height = max(20, approx_lines * 17)
            else:
                cell.font = body_font
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                approx_lines = max(1, len(line_clean) // 120 + 1)
                ws.row_dimensions[current_row].height = max(20, approx_lines * 17)

            current_row += 1

        # Baris kosong penutup
        ws.row_dimensions[current_row].height = 15

        # Atur lebar kolom agar seluruh teks pas dan rapi
        ws.column_dimensions['A'].width = 3
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 16
        ws.column_dimensions['D'].width = 14
        ws.column_dimensions['E'].width = 18
        ws.column_dimensions['F'].width = 16
        ws.column_dimensions['G'].width = 16
        ws.column_dimensions['H'].width = 16
        ws.column_dimensions['I'].width = 14
        ws.column_dimensions['J'].width = 14
        ws.column_dimensions['K'].width = 12
        ws.column_dimensions['L'].width = 12

        wb.save(self.excel_path)
        return self.excel_path

    def sync_to_google_sheets(self, creds_path: str, spreadsheet_id: str, results: List[Dict[str, Any]], metrics: Dict[str, Any], summary_text: str):
        """
        Syncs data directly to a live Google Sheet using gspread.
        """
        if not creds_path or not os.path.exists(creds_path) or not spreadsheet_id:
            print("[Info] Google Sheets credentials or Spreadsheet ID not specified. Skipping live sync.")
            return False

        try:
            import gspread
            gc = gspread.service_account(filename=creds_path)
            sh = gc.open_by_key(spreadsheet_id)

            # Update or create 'Working_Paper_Result'
            try:
                ws_result = sh.worksheet("Working_Paper_Result")
            except Exception:
                ws_result = sh.add_worksheet(title="Working_Paper_Result", rows=50, cols=15)

            # Prepare data matrix
            rows_data = [
                ["Date", "Voucher No", "Description", "Amount", "Realization Date", "Realization No. Voucher", "Realization Amount", "Saldo", "Status"]
            ]
            for r in results:
                rows_data.append([
                    str(r['date'])[:10] if r['date'] else '',
                    r['advance_voucher'],
                    r['desc'],
                    r['amount'],
                    r['realization_date'],
                    r['realization_voucher'],
                    r['realization_amount'],
                    r['saldo'],
                    r['status']
                ])

            ws_result.clear()
            ws_result.update(rows_data, 'A1')

            # Update or create 'Dashboard'
            try:
                ws_dash = sh.worksheet("Dashboard")
            except Exception:
                ws_dash = sh.add_worksheet(title="Dashboard", rows=60, cols=10)

            dash_data = [
                ["SOUTHCITY ADVANCE SETTLEMENT DASHBOARD"],
                [f"Settlement Rate: {metrics['settlement_rate_pct']}%"],
                [],
                ["METRIC", "VALUE"],
                ["Total Advance", metrics['total_advance']],
                ["Total Realisasi", metrics['total_realization']],
                ["Sisa Saldo", metrics['total_saldo']],
                ["Items Settled", f"{metrics['count_settled']} / {metrics['total_items']}"],
                [],
                ["EXECUTIVE SUMMARY"],
            ]
            for line in summary_text.split("\n"):
                if line.strip():
                    dash_data.append([line.strip()])

            ws_dash.clear()
            ws_dash.update(dash_data, 'A1')
            print(f"[Success] Synced successfully to Google Sheets URL: https://docs.google.com/spreadsheets/d/{spreadsheet_id}")
            return True
        except Exception as e:
            print(f"[Error] Failed to sync to Google Sheets: {e}")
            return False
